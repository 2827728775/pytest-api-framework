"""数据驱动加载器：支持 YAML 与 Excel 两种用例来源。"""
import os

import yaml


def load_yaml(path: str):
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def load_excel(path: str, sheet: str = "Sheet1"):
    """Excel 用例：第一行为表头，后续每行转成 dict。需安装 openpyxl。"""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("使用 Excel 数据驱动请先安装：pip install openpyxl") from e

    wb = load_workbook(path)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = rows[0]
    return [
        dict(zip(header, row))
        for row in rows[1:]
        if any(cell is not None for cell in row)
    ]


def load(path: str, sheet: str = "Sheet1"):
    """按扩展名自动选择加载方式。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".yaml", ".yml"):
        return load_yaml(path)
    if ext in (".xlsx", ".xls"):
        return load_excel(path, sheet)
    raise ValueError(f"不支持的数据文件类型：{ext}")
